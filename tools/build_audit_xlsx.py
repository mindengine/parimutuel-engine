"""Build the grade-audit workbook from a rules-graded DB and Grant's breakdown DB.

usage: PYTHONPATH=. python tools/build_audit_xlsx.py RULES_DB BREAKDOWN_DB OUT.xlsx "Oct 25, 2026"
"""
import sqlite3
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pe import analytics

rules_db, breakdown_db, out, day = sys.argv[1:5]
con = sqlite3.connect(rules_db); con.row_factory = sqlite3.Row
con2 = sqlite3.connect(breakdown_db); con2.row_factory = sqlite3.Row

F = Font(name="Arial", size=10)
B = Font(name="Arial", size=10, bold=True, color="FFFFFF")
H = PatternFill("solid", fgColor="1F3864")


def sheet(ws, df, widths=None):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font = B; c.fill = H; c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.font = F
    for i, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, max(9, min(40, len(str(c)) + 3)))
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions


wb = Workbook()
run = pd.read_sql_query("""select c.track Track, g.race Race, r.pace_scenario Pace, r.qsp_sum "QSP sum", g.pgm Pgm,
 g.horse_name Horse, g.bucket Grade, g.modifiers Mods, g.ml_odds ML, o.finish_pos Finish,
 round(o.final_odds,2) "Final odds", g.tags Why
 from pre_race_grades g join cards c using(card_id) join races r using(card_id,race)
 join post_race_official o on o.track=c.track and o.race_date=c.race_date and o.race=g.race and o.pgm=g.pgm
 where o.finish_pos is not null order by c.track, g.race, o.finish_pos""", con)
run["Grade group"] = run.Grade.map(lambda b: "F / F-" if b in ("F", "F-") else b)
run["$2 win return"] = [f"=IF(J{i}=1,2*(K{i}+1),0)" for i in range(2, len(run) + 2)]
n = len(run) + 1

ws = wb.active; ws.title = "Summary"
tracks = ", ".join(sorted(run.Track.unique()))
ws.append([f"Grade performance: every starter, {day} ({tracks})"]); ws["A1"].font = Font(name="Arial", size=12, bold=True)
ws.append(["Grades from the rules grader (pe/grading.py) run on Grant's distill.py output. "
           "59% exact agreement with the 204 horses Grant graded himself on Oct 25."])
ws["A2"].font = Font(name="Arial", size=9, italic=True)
ws.append([]); ws.append(["Grade", "Starters", "Wins", "Win %", "Top 3", "Top 3 %", "$2 win ROI"])
for c in ws[4]:
    c.font = B; c.fill = H
R = f"Runners!$M$2:$M${n}"; Fn = f"Runners!$J$2:$J${n}"; Ret = f"Runners!$N$2:$N${n}"
for i, g in enumerate(["A+", "A", "B", "C", "D", "F / F-"]):
    r = 5 + i
    ws.append([g, f"=COUNTIF({R},A{r})", f"=COUNTIFS({R},A{r},{Fn},1)", f"=IF(B{r}=0,0,C{r}/B{r})",
               f'=COUNTIFS({R},A{r},{Fn},"<=3")', f"=IF(B{r}=0,0,E{r}/B{r})",
               f"=IF(B{r}=0,0,SUMIF({R},A{r},{Ret})/(2*B{r})-1)"])
ws.append(["Total", "=SUM(B5:B10)", "=SUM(C5:C10)", "=IF(B11=0,0,C11/B11)", "=SUM(E5:E10)",
           "=IF(B11=0,0,E11/B11)", f"=IF(B11=0,0,SUM({Ret})/(2*B11)-1)"])
for row in ws.iter_rows(min_row=5, max_row=11):
    for c in row:
        c.font = Font(name="Arial", size=10, bold=(row[0].row == 11))
    for k in (3, 5, 6):
        row[k].number_format = "0.0%;(0.0%);-"
ws.append([]); ws.append(["Key findings"]); ws.cell(ws.max_row, 1).font = Font(name="Arial", size=11, bold=True)
for t in ["All 49 winners now have a grade (none ungraded).",
          "A+ and A won 21 of 49 races from 92 starters; A+ alone was 3 for 6.",
          "A horses won 12 of 42 in Honest-pace races but only 2 of 24 in Lone Speed races.",
          "C (meltdown closers) won 1 of 11; D won 3 of 58.",
          "The only rules F- (Mo for the King) won. Grant's retail-trap F grades need public-pick data that the distill file doesn't have.",
          "Five cards is far too small a sample to judge the edge."]:
    ws.append([t]); ws.cell(ws.max_row, 1).font = F
ws.column_dimensions["A"].width = 14
for col in "BCDEFG":
    ws.column_dimensions[col].width = 11

sheet(wb.create_sheet("Runners"), run, {"Horse": 22, "Why": 45, "Pace": 12})

ag = pd.read_sql_query("""select substr(r.card_id,1,instr(r.card_id,'-')-1) Track, r.race Race, r.pgm Pgm,
 r.horse_name Horse, g.bucket "Grant grade", r.bucket "Rules grade",
 case when g.bucket=r.bucket then 'Match' else 'Differs' end Result
 from pre_race_grades r join pre_race_grades g on g.card_id=substr(r.card_id,1,length(r.card_id)-9)
 and g.race=r.race and g.pgm=r.pgm where r.card_id like '%rules_v1' and g.bucket is not null order by 1,2,3""", con2)
sheet(wb.create_sheet("Rules vs Grant"), ag, {"Horse": 22})

g = analytics.grade_tickets(con2); g = g[~g.card_id.str.contains("rules")]
t = pd.DataFrame({"Card": g.card_id.str[:-11], "Pool": g.pool, "First race": g.race, "Cost ($)": g.cost,
                  "Official result": g.result,
                  "Status": g.hit.map(lambda h: "Not offered" if h is None or pd.isna(h) else ("Hit" if h else "Lost")),
                  "Returned ($)": g.returned, "Label": g.label.str.strip("* ").str[:70]})
ws = wb.create_sheet("Tickets"); sheet(ws, t, {"Official result": 20, "Label": 60}); m = len(t) + 1
for r in range(2, m + 1):
    ws[f"D{r}"].number_format = ws[f"G{r}"].number_format = "$#,##0.00"
ws.append([]); ws.append(["Playable cost", "", "", f'=SUMIF(F2:F{m},"<>Not offered",D2:D{m})'])
ws.append(["Returned", "", "", f"=SUM(G2:G{m})"]); ws.append(["Net", "", "", f"=D{m+3}-D{m+2}"])
for r in range(m + 2, m + 5):
    ws[f"A{r}"].font = Font(name="Arial", size=10, bold=True); ws[f"D{r}"].number_format = "$#,##0.00;($#,##0.00)"
ws.append([])
ws.append(["Caveat: CD and GP backups were trimmed in prose; the full grid cost is counted here. "
           "The winning PRX ticket was 'scaled to 10 units', so its real stake is unclear."])
ws.cell(ws.max_row, 1).font = Font(name="Arial", size=9, italic=True)

l = analytics.lott_audit(con); l = l[l.system_covered].copy()
lo = pd.DataFrame({"Card": l.card_id.str.split("-").str[0], "Pool": l.pool, "Closing race": l.close_race,
                   "Winning combo": l.winning_combo, "Grades": l.buckets, "Payout ($)": l.payout, "Base ($)": l.base})
ws = wb.create_sheet("Covered Payouts"); sheet(ws, lo)
for r in range(2, len(lo) + 2):
    ws[f"F{r}"].number_format = ws[f"G{r}"].number_format = "$#,##0.00"

ws = wb.create_sheet("Notes")
for x in ["Sources: Brisnet PP Data Files run through Grant's distill.py; official results and payouts from Brisnet "
          "Comprehensive Charts; tickets and Grant's own grades from his game-theory breakdowns.",
          "Grades are a deterministic approximation of Grant's system built from the dossier's bucket definitions and "
          "tuned only against his own grades, never against results.",
          "Scratches removed using the official charts. Pace scenario: QSP sum <=15 Lone Speed, 16-29 Honest, >=30 Meltdown.",
          "Covered Payouts lists official payouts whose winners were all graded A+ through D. Tickets come from Grant's "
          "breakdowns, which don't include his single-race bets.",
          "$2 win ROI uses final odds as the payout.",
          "Generated by mindengine/parimutuel-engine (tools/build_audit_xlsx.py)."]:
    ws.append([x]); ws.cell(ws.max_row, 1).font = F
ws.column_dimensions["A"].width = 120
wb.save(out)
print("saved", out)
